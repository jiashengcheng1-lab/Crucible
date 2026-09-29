"""The model's own state as context: the three statements as mapped now, the forecast and DCF as last built, and the
workbook sheets. Follow-up questions are answered against this alongside the filings, so "what is under the balance
sheet now?" is answerable, and the UI previews the same tables."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd

from .schema import ITEM_BY_KEY, STANDARD_ITEMS

STATEMENT_NAMES = {"IS": "income statement", "BS": "balance sheet", "CF": "cash flow statement"}


def statements_state(data_dir: Path, ticker: str, as_of: str | None = None) -> dict:
    """The mapped history (USD m) with the source of every line and the mapping notes."""
    from .mapping import load_history

    hist, reports = load_history(Path(data_dir), ticker, as_of=as_of)
    return {"hist": hist, "reports": reports, "years": [int(y) for y in hist.columns]}


def _fmt(v) -> str:
    try:
        return f"{float(v):,.1f}"
    except (TypeError, ValueError):
        return "n/a"


def statement_items(state: dict, ticker: str) -> list[dict]:
    """One context item per mapped line ("cash [cash]: FY2023 29,965; ... from facts: CashAndCashEquivalents...") plus one
    per statement for the residual lines and the mapping notes."""
    hist, reports = state["hist"], state["reports"]
    out = []
    for st in ("IS", "BS", "CF"):
        rep = reports.get(st)
        name = STATEMENT_NAMES[st]
        src_tag = f"{name} as mapped now (data/{ticker.upper()})"
        for it in [i for i in STANDARD_ITEMS if i.statement == st]:
            if it.key not in hist.index:
                continue
            vals = "; ".join(f"FY{y} {_fmt(hist.at[it.key, y])}" for y in hist.columns)
            src = (rep.mapped.get(it.key) if rep else None) or "derived or filled"
            out.append({"source": src_tag, "text": f"{it.label} [{it.key}] (USD m): {vals}; mapped from {src}", "kind": "state", "statement": st})
        resid = [k for k in hist.index if k.startswith("other_") or k in ("d_nwc", "net_debt_issuance", "nci_and_other", "residual")]
        resid = [k for k in resid if k in hist.index and k not in {i.key for i in STANDARD_ITEMS}]
        if resid:
            rv = "; ".join(f"{k}: " + ", ".join(f"FY{y} {_fmt(hist.at[k, y])}" for y in hist.columns) for k in resid)
            out.append({"source": src_tag, "text": f"{name} residual lines (what the unmapped facts fall into): {rv}", "kind": "state", "statement": st})
        if rep is not None:
            if rep.residuals:
                out.append({"source": src_tag, "text": f"{name} residual size (max share of base): " + "; ".join(f"{k} {v:.1%}" for k, v in rep.residuals.items()), "kind": "state", "statement": st})
            if rep.unmapped:
                out.append({"source": src_tag, "text": f"{name} labels with numbers that map to nothing yet: " + "; ".join(rep.unmapped[:20]), "kind": "state", "statement": st})
            for n in rep.notes[:8]:
                out.append({"source": src_tag, "text": f"{name} mapping note: {n[:400]}", "kind": "state", "statement": st})
    return out


def model_items(ticker: str, outputs: Path = Path("outputs")) -> list[dict]:
    """The forecast and DCF as last built (outputs/<T>_forecast.csv, outputs/<T>_model_summary.json) and the workbook's sheets."""
    t = ticker.upper()
    o = Path(outputs)
    out = []
    fc_path, sum_path, wb_path = o / f"{t}_forecast.csv", o / f"{t}_model_summary.json", o / f"{t}_model.xlsx"
    if fc_path.exists():
        fc = pd.read_csv(fc_path, index_col=0)
        fc.columns = [str(c) for c in fc.columns]
        src = f"model as last built (outputs/{t}_model.xlsx)"
        for key in ("revenue", "cogs", "gross_profit", "sga", "operating_income", "net_income", "d_and_a", "capex", "d_nwc", "cfo", "cash", "total_assets",
                    "total_liabilities", "total_equity", "short_term_debt", "long_term_debt", "shares_diluted"):
            if key in fc.index:
                label = ITEM_BY_KEY[key].label if key in ITEM_BY_KEY else key
                out.append({"source": src, "text": f"{label} [{key}] history and forecast (USD m): " + "; ".join(f"FY{c} {_fmt(fc.at[key, c])}" for c in fc.columns), "kind": "state", "statement": "MODEL"})
    if sum_path.exists():
        try:
            s = json.loads(sum_path.read_text())
            d = s.get("dcf", {})
            out.append({"source": f"model summary as last built (outputs/{t}_model.xlsx)", "kind": "state", "statement": "MODEL",
                        "text": (f"DCF: WACC {d.get('wacc')}, terminal growth {d.get('terminal_growth')}, enterprise value {_fmt(d.get('enterprise_value'))}, net debt {_fmt(d.get('net_debt'))}, "
                                 f"equity value {_fmt(d.get('equity_value'))}, diluted shares {_fmt(d.get('diluted_shares'))}m, per share {d.get('per_share')}, terminal share of EV {d.get('tv_share_of_ev')}; "
                                 f"scenarios per share {s.get('scenarios')}; checks {s.get('checks', {}).get('ok')}; driver notes {s.get('driver_notes')}; warnings {s.get('warnings')}")})
            if s.get("drivers"):
                out.append({"source": f"model drivers as last built (Inputs sheet of outputs/{t}_model.xlsx)", "kind": "state", "statement": "MODEL",
                            "text": "forecast drivers by year: " + "; ".join(f"{k} {v}" for k, v in s["drivers"].items())})
        except Exception:
            pass
    if wb_path.exists():
        try:
            from openpyxl import load_workbook

            wb = load_workbook(wb_path, read_only=True)
            out.append({"source": f"workbook outputs/{t}_model.xlsx", "kind": "state", "statement": "MODEL",
                        "text": "Excel sheets in the model workbook: " + ", ".join(wb.sheetnames) + ". Inputs holds the driver assumptions (blue cells), Model the three statements "
                                "with formulas, DCF the valuation, Scenarios the bull/base/bear driver sets with the selector in Inputs!B1, Elasticity the per-driver sensitivities, Sources the filings used."})
            wb.close()
        except Exception:
            pass
    return out


def state_items(data_dir: Path, ticker: str, outputs: Path = Path("outputs")) -> list[dict]:
    """Everything the model currently holds, as context items (statements first, then the model)."""
    items: list[dict] = []
    try:
        items += statement_items(statements_state(data_dir, ticker), ticker)
    except Exception as e:
        items.append({"source": "model state", "text": f"the three statements could not be loaded: {e}", "kind": "state", "statement": "IS"})
    items += model_items(ticker, outputs)
    return items


def sheet_frame(path: Path, sheet: str, max_rows: int = 200, max_cols: int = 30) -> pd.DataFrame:
    """A workbook sheet as a table for preview: values as stored, formulas shown as their text."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    ws = wb[sheet]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            break
        rows.append([("" if v is None else v) for v in row[:max_cols]])
    wb.close()
    if not rows:
        return pd.DataFrame()
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    df = pd.DataFrame(rows, columns=[f"{chr(65 + j)}" if j < 26 else f"C{j}" for j in range(width)])
    return df.astype(str)  # mixed numbers, text and formula strings in one column: keep the preview as text


def sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names

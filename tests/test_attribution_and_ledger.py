import numpy as np

from crucible.attribution import attribution_from_logs, attribution_random, design_random_subsets, fit_logistic, fit_ols, group_ids
from crucible.evidence import build_packet
from crucible.harness import DebateCache
from crucible.ledger import MappingLedger
from crucible.llm import MockLLM
from crucible.synthetic import synthetic_history


def test_random_design_varies_every_optional_group():
    d = design_random_subsets(["target history", "a", "b", "c"], k=8, seed=3)
    assert all(r["target history"] == 1 for r in d)
    for g in ("a", "b", "c"):
        col = [r[g] for r in d]
        assert 0 < sum(col) < len(col)


def test_ols_recovers_a_planted_effect():
    rng = np.random.default_rng(0)
    X = rng.integers(0, 2, (60, 3)).astype(float)
    y = 0.10 + 0.05 * X[:, 0] - 0.02 * X[:, 1] + rng.normal(0, 0.005, 60)
    fit = fit_ols(X, y, ["a", "b", "c"])
    eff = {r["group"]: r["effect"] for r in fit["rows"]}
    assert abs(eff["a"] - 0.05) < 0.01 and abs(eff["b"] + 0.02) < 0.01 and abs(eff["c"]) < 0.01 and fit["r2"] > 0.9
    lg = fit_logistic(X, (y > 0.12).astype(float), ["a", "b", "c"])
    assert lg["rows"][0]["log_odds"] > 0


def test_attribution_random_with_mock(tmp_path):
    pk = build_packet("revenue_growth_3y", "Test Co", synthetic_history())
    df, fit, results = attribution_random(pk, MockLLM(noise=0.0), k=8, workers=2, cache=DebateCache(tmp_path / "c"), progress=None)
    assert fit["n"] == 8 and len(results) == 8 and set(df.group) == set(group_ids(pk))
    logs = [{"assumption": "revenue_growth_3y", "company": "Test Co", "evidence_ids": r.evidence_ids, "verdict": r.verdict.model_dump()} for r in results]
    fl = attribution_from_logs(logs, pk)
    assert fl["n_runs"] == 8 and "ols_base" in fl


def test_ledger_roundtrip_and_excel(tmp_path):
    L = MappingLedger(tmp_path)
    L.upsert("XYZ", "IS", "label", "Cost of goods", "cogs", 0.6, [{"target": "other_opex", "confidence": 0.3}], "accepted", "rule:label", evidence="FY2025: 100")
    L.upsert("XYZ", "IS", "label", "Weird line", "residual", 0.4, [], "pending", "llm:mock")
    L.save()
    L2 = MappingLedger(tmp_path)
    assert L2.accepted("XYZ", "IS", "label") == {"Cost of goods": "cogs"} and len(L2.rows("XYZ", status="pending")) == 1
    L2.upsert("XYZ", "IS", "label", "Cost of goods", "sga", 0.6, [], "accepted", "rule:label")  # no decision yet: rule can overwrite
    assert L2.accepted("XYZ", "IS", "label")["Cost of goods"] == "sga"
    L2.decide("XYZ", "Cost of goods", "accepted", target="cogs", decided_by="analyst", statement="IS")
    L2.upsert("XYZ", "IS", "label", "Cost of goods", "sga", 0.6, [], "accepted", "rule:label")  # analyst decided: rule must not overwrite
    assert L2.accepted("XYZ", "IS", "label")["Cost of goods"] == "cogs"
    x = L2.to_excel(tmp_path / "l.xlsx", "XYZ")
    from openpyxl import load_workbook
    wb = load_workbook(x); ws = wb["mapping_ledger"]
    hdr = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[hdr.index("source")].value == "Weird line":
            row[hdr.index("status")].value = "accepted"; row[hdr.index("target")].value = "other_opex"
    wb.save(x)
    assert L2.from_excel(x) == 1 and L2.accepted("XYZ", "IS", "label")["Weird line"] == "other_opex"


def test_rule_rows_do_not_feed_back_only_decisions_do(tmp_path):
    L = MappingLedger(tmp_path)
    L.upsert("XYZ", "IS", "concept", "us-gaap:RevenueFromContractWithCustomerExcludingAssessedTax", "revenue", 0.95, [], "accepted", "rule:concept")
    L.upsert("XYZ", "IS", "concept", "us-gaap:Revenues", "revenue", 0.95, [], "accepted", "rule:concept")
    assert L.decided_links("XYZ", "IS", "concept") == {}          # rule records never reorder rule priority
    L.decide("XYZ", "us-gaap:Revenues", "accepted", target="revenue", decided_by="analyst", statement="IS")
    assert L.decided_links("XYZ", "IS", "concept") == {"us-gaap:Revenues": "revenue"}
    x = L.to_excel(tmp_path / "v.xlsx", "XYZ")
    from openpyxl import load_workbook
    wb = load_workbook(x)
    assert "lists" in wb.sheetnames and wb["lists"].sheet_state == "hidden"
    dv = wb["mapping_ledger"].data_validations.dataValidation
    assert dv and all(len(d.formula1) < 255 for d in dv)

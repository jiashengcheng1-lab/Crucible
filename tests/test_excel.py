import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from crucible.excel import MODEL_ROWS, write_model
from crucible.model import Drivers, WACCInputs, dcf, forecast
from crucible.synthetic import synthetic_history

RECALC = Path("/mnt/skills/public/xlsx/scripts/recalc.py")


@pytest.mark.skipif(not RECALC.exists() or shutil.which("soffice") is None and not Path("/usr/bin/libreoffice").exists(),
                    reason="LibreOffice recalc not available")
def test_workbook_recalculates_and_matches_engine(tmp_path):
    hist = synthetic_history()
    drv = Drivers.from_history(hist, years=5)
    w = WACCInputs()
    out = write_model(hist, drv, w, 0.025, tmp_path / "m.xlsx", company="Test Co", mid_year=True)
    res = subprocess.run([sys.executable, str(RECALC), str(out), "90"], capture_output=True, text=True, check=True)
    rep = json.loads(res.stdout)
    assert rep["status"] == "success" and rep["total_errors"] == 0, rep
    fc = forecast(hist, drv)
    r = dcf(fc, drv, w.wacc, 0.025, mid_year=True)
    wb = load_workbook(out, data_only=True)
    ws, wd = wb["Model"], wb["DCF"]
    rows = {k: i + 4 for i, (k, _, _) in enumerate(MODEL_ROWS)}
    years = list(hist.columns) + drv.years
    col = {y: get_column_letter(3 + i) for i, y in enumerate(years)}
    for key in ("revenue", "net_income", "cash", "total_assets", "balance_check", "cfo"):
        for y in drv.years:
            assert ws[f"{col[y]}{rows[key]}"].value == pytest.approx(fc.loc[key, y], abs=1e-6)
    assert wd["B11"].value == pytest.approx(w.wacc)
    assert wd["B34"].value == pytest.approx(r.per_share, rel=1e-9)
    assert wd["E41"].value == pytest.approx(r.per_share, rel=1e-9)  # sensitivity centre cell

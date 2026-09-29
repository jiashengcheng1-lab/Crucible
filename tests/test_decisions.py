"""Decisions, review, export, brief, scenarios, quarterly history, dilution: all offline with the mock model."""
import io
from pathlib import Path
import json
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from crucible.brief import brief_markdown, write_brief
from crucible.decisions import DecisionLedger, DecisionSpec, Option, decide, export_workbook, review_terminal
from crucible.dossier import DOSSIER_SPEC, FAMILIES, allocation_items, competitor_names, decision_specs, family_items, kpi_items
from crucible.evidence import EvidenceItem, EvidencePacket
from crucible.facts import quarterly_from_facts
from crucible.inputs import _dilution
from crucible.ledger import LEDGER_COLUMNS, MappingLedger
from crucible.llm import MockLLM
from crucible.model import Drivers, dcf, elasticity, forecast, scenario_values, shifted_drivers
from crucible.synthetic import synthetic_history


def _packet():
    items = [EvidenceItem(id="E1", kind="historical_metric", source="10-K FY2025", content="Test Co revenue growth 20.0% FY2025", value=0.2, unit="pct"),
             EvidenceItem(id="E2", kind="filing_text", source="10-K Item 1 [segments]", content='"We report two segments: Products and Services, with Products at 70 percent of revenue."'),
             EvidenceItem(id="E3", kind="filing_text", source="10-K Item 7 [kpi]", content='"Backlog was 7.5 billion at year end and orders increased 30 percent."'),
             EvidenceItem(id="E4", kind="macro", source="FRED INDPRO", content="industrial production, y/y: +1.2% as of 2026-08-01 (FRED INDPRO)", value=1.2, unit="pct")]
    return EvidencePacket(company="Test Co", assumption=DOSSIER_SPEC, items=items, as_of="2026-09-24")


def _spec(n=4):
    opts = [Option("reported segments", "reported", "as disclosed"), Option("geography", "geography", "by region"),
            Option("product lines", "product", "by product"), Option("single segment", "single", "one line")][:n]
    return DecisionSpec("segment_scheme", "Which segment scheme?", opts, considerations=("disclosure",), settles="the segment note")


def test_tournament_ranks_every_option_and_caches(tmp_path):
    pk, spec = _packet(), _spec()
    res = decide(pk, spec, MockLLM(), seed=1, cache_dir=tmp_path)
    labels = [v.label for v in res.ranking]
    assert len(labels) == 4 and set(labels) == {o.label for o in spec.options}
    assert res.ranking[0].label == "reported segments" and res.ranking[0].citations and res.ranking[0].questions
    assert len(res.pairings) == 3 and len(list(tmp_path.glob("decision_*.json"))) == 3
    n_calls = len(res.call_stats)
    res2 = decide(pk, spec, MockLLM(), seed=1, cache_dir=tmp_path)
    assert len(res2.call_stats) == 0 and [v.label for v in res2.ranking] == labels  # second run served from cache
    cheap = decide(pk, spec, MockLLM(), seed=1, cheap=True)
    assert len(cheap.call_stats) == 1 and len(cheap.ranking) == 4 and n_calls == 9
    assert all(v.confidence > 0 for v in cheap.ranking) and cheap.ranking[0].label == "reported segments"  # every option matched and ranked
    spec_t = DecisionSpec("market_cycle", "Where?", spec.options, tournament=True)
    cheap_t = decide(pk, spec_t, MockLLM(), seed=1, cheap=True)
    assert len(cheap_t.pairings) == 3 and len(cheap_t.call_stats) == 3  # tournament kept under --cheap: judge only, no advocates


def test_two_option_debate_records_arguments_with_verified_quotes(tmp_path):
    pk = _packet()
    spec = DecisionSpec("pricing_power", "High or low?", [Option("high", "high", "can raise prices"), Option("low", "low", "price taker")])
    res = decide(pk, spec, MockLLM(), seed=3)
    assert len(res.pairings) == 1 and res.ranking[0].label in ("high", "low")
    args = res.pairings[0]["arguments"]
    assert args["A"]["claims"] and all(c["quote"] for c in args["A"]["claims"])


def test_ledger_review_and_export(tmp_path):
    pk, spec = _packet(), _spec()
    res = decide(pk, spec, MockLLM(), seed=1)
    led = DecisionLedger(tmp_path)
    assert led.record("TST", res, pk, spec) == 4
    led.save()
    led2 = DecisionLedger(tmp_path)
    rows = led2.rows("TST", "segment_scheme")
    assert len(rows) == 4 and set(rows.status) == {"pending"} and rows.iloc[0].citation.startswith("10-K")
    assert rows.iloc[0].questions and rows.iloc[0].duplicates  # questions and considerations recorded
    # non-interactive review: accept the winner, mark one unsure, reject one, leave one pending
    counts = review_terminal(led2, "TST", decisions={"reported segments": ("accepted", None, "matches the note"), "geography": ("unsure", None, "ask IR"),
                                                     "product lines": ("rejected", None, "")})
    assert counts == {"accepted": 1, "rejected": 1, "unsure": 1, "skipped": 1, "auto": 0}
    assert led2.selected("TST", "segment_scheme") == ["reported"]
    # interactive review over a fake terminal: note then unsure on the remaining pending row
    out = io.StringIO()
    counts2 = review_terminal(led2, "TST", io_in=io.StringIO("nneeds the revenue note\nu\n"), io_out=out)
    assert counts2["unsure"] == 1 and "[a]ccept" in out.getvalue()
    unsure = led2.rows("TST", status="unsure")
    assert len(unsure) == 2 and "needs the revenue note" in " ".join(unsure.note)
    # export: final workbook without unsure rows, unsure workbook with them, brief tab from a mock brief
    brief = write_brief(pk, MockLLM())
    final, unsure_path = export_workbook(tmp_path / "TST_final.xlsx", "TST", MappingLedger(tmp_path), led2, brief)
    wb = load_workbook(final)
    assert wb.sheetnames == ["mapping_ledger", "decision_ledger", "brief"]
    ws = wb["decision_ledger"]
    assert [c.value for c in ws[1]] == LEDGER_COLUMNS and ws.max_row == 3  # header + accepted + rejected
    assert wb["brief"]["A2"].value == "what_it_does" and "E1" in wb["brief"]["C2"].value
    wu = load_workbook(unsure_path)
    assert wu["unsure"].max_row == 3 and "# Test Co brief" in brief_markdown(brief)


def test_dossier_builders_and_specs():
    hist = synthetic_history()
    sections = {"Item 1 Business": "We compete with Alpha Corp, Beta Industries and Gamma Systems in the thermal market. Our segments are Products and Services.",
                "Competition excerpt": "Competitors include Alpha Corp, Beta Industries and Gamma Systems.",
                "Item 7 MD&A": "Backlog grew to $7.5 billion. Our hashrate reached 50 EH/s during the year. Pricing power was strong."}
    names = competitor_names(sections, "Synthetic Co")
    assert names[:3] == ["Alpha Corp", "Beta Industries", "Gamma Systems"]
    fam = family_items(sections, "kpi", "10-K", 1)
    assert fam and fam[0].kind == "filing_text" and "[kpi]" in fam[0].source
    rel = [{"filing_date": "2026-02-10", "text": "Fourth quarter results. Backlog was $7.5 billion at year end. Orders increased 30% year over year."}]
    ks = kpi_items(rel, "Synthetic Co", 10)
    assert {i.source.split("[kpi:")[1].rstrip("]") for i in ks} >= {"backlog", "orders"}
    alloc = allocation_items(hist, "Synthetic Co", 20)
    assert len(alloc) == 2 and "capital allocation" in alloc[0].content and "ROIC proxy" in alloc[1].content
    specs = decision_specs(sections, ["ETN"], "Synthetic Co")
    keys = [s.key for s in specs]
    assert keys[:4] == ["segment_scheme", "unit_economics", "kpi_scheme", "peer_set"] and "market_cycle" in keys and "dcf_method" in keys
    peer = next(s for s in specs if s.key == "peer_set")
    assert peer.options[0].value.startswith("named:Alpha Corp") and peer.options[1].value == "current:ETN"
    assert all(s.options and s.question and s.settles for s in specs) and set(FAMILIES) >= {"segments", "kpi", "incentives"}


def test_scenarios_elasticity_and_nwc_method():
    hist = synthetic_history()
    drv = Drivers.from_history(hist, years=5)
    fc = forecast(hist, drv)
    base = dcf(fc, drv, 0.10, 0.025).as_dict()["per_share"]
    scn = scenario_values(hist, drv, 0.10, 0.025, {"base": {"prob": 0.5}, "bull": {"growth_delta": 0.05, "margin_delta": 0.01, "prob": 0.25},
                                                    "bear": {"growth_delta": -0.05, "margin_delta": -0.01, "wacc_delta": 0.005, "prob": 0.25}})
    assert scn["bull"]["per_share"] > scn["base"]["per_share"] > scn["bear"]["per_share"]
    assert abs(scn["base"]["per_share"] - base) < 1e-9 and "_probability_weighted" in scn
    el = elasticity(hist, drv, 0.10, 0.025)
    assert {r["driver"] for r in el} >= {"revenue_growth", "gross_margin", "wacc", "terminal_growth", "dso"}
    growth = next(r for r in el if r["driver"] == "revenue_growth")
    assert growth["delta"] > 0 and next(r for r in el if r["driver"] == "wacc")["delta"] < 0
    d2 = shifted_drivers(drv, growth_delta=0.02, taper=0.5)
    assert abs((d2.revenue_growth[0] - drv.revenue_growth[0]) - 0.02) < 1e-12 and abs((d2.revenue_growth[1] - drv.revenue_growth[1]) - 0.01) < 1e-12
    drv.nwc_method = "pct_revenue"
    fc2 = forecast(hist, drv)
    last_rev, last_ar = float(hist.loc["revenue"].iloc[-1]), float(hist.loc["receivables"].iloc[-1])
    y = drv.years[0]
    assert abs(fc2.loc["receivables", y] - fc2.loc["revenue", y] * last_ar / last_rev) < 1e-6
    assert abs(fc2.loc["balance_check", y]) < 1e-6  # the balance sheet still ties under the other working capital method
    drv.shares_override = 123.0
    assert abs(dcf(fc2, drv, 0.10, 0.025).as_dict()["diluted_shares"] - 123.0) < 1e-9


def _facts_frame():
    rows = []
    def add(concept, start, end, value, form, filing):
        rows.append({"concept": concept, "period_start": pd.Timestamp(start), "period_end": pd.Timestamp(end), "numeric_value": value, "form_type": form,
                     "filing_date": pd.Timestamp(filing), "period_type": "duration", "fiscal_period": "", "fiscal_year": pd.Timestamp(end).year, "accession": "x",
                     "unit": "USD", "dimension": None})
    for fy in (2024, 2025):
        q = [(f"{fy}-01-01", f"{fy}-03-31", 100 + fy - 2024), (f"{fy}-04-01", f"{fy}-06-30", 110), (f"{fy}-07-01", f"{fy}-09-30", 120)]
        for st, en, v in q:
            add("us-gaap:Revenues", st, en, v * 1e6, "10-Q", pd.Timestamp(en) + pd.Timedelta(days=40))
            add("us-gaap:CostOfRevenue", st, en, v * 0.6e6, "10-Q", pd.Timestamp(en) + pd.Timedelta(days=40))
        add("us-gaap:Revenues", f"{fy}-01-01", f"{fy}-12-31", 500e6, "10-K", f"{fy + 1}-02-20")
        add("us-gaap:CostOfRevenue", f"{fy}-01-01", f"{fy}-12-31", 300e6, "10-K", f"{fy + 1}-02-20")
    return pd.DataFrame(rows)


def test_quarterly_q4_is_year_minus_nine_months():
    q = quarterly_from_facts(_facts_frame())
    assert list(q.columns)[:4] == ["FY2024Q1", "FY2024Q2", "FY2024Q3", "FY2024Q4"]
    assert abs(q.loc["revenue", "FY2025Q4"] - (500 - (101 + 110 + 120))) < 1e-6
    assert abs(q.loc["gross_profit", "FY2025Q1"] - (101 - 101 * 0.6)) < 1e-6  # derived when GrossProfit is untagged
    q2 = quarterly_from_facts(_facts_frame(), as_of="2025-06-01")
    assert "FY2025Q4" not in q2.columns  # the FY2025 10-K is not filed yet at that date


def test_treasury_stock_dilution():
    f = pd.DataFrame([
        {"concept": "us-gaap:ShareBasedCompensationArrangementByShareBasedPaymentAwardOptionsOutstandingNumber", "numeric_value": 10e6, "period_end": pd.Timestamp("2025-12-31"), "filing_date": pd.Timestamp("2026-02-20"), "form_type": "10-K"},
        {"concept": "us-gaap:ShareBasedCompensationArrangementByShareBasedPaymentAwardOptionsOutstandingWeightedAverageExercisePrice", "numeric_value": 20.0, "period_end": pd.Timestamp("2025-12-31"), "filing_date": pd.Timestamp("2026-02-20"), "form_type": "10-K"},
        {"concept": "us-gaap:ShareBasedCompensationArrangementByShareBasedPaymentAwardEquityInstrumentsOtherThanOptionsNonvestedNumber", "numeric_value": 5e6, "period_end": pd.Timestamp("2025-12-31"), "filing_date": pd.Timestamp("2026-02-20"), "form_type": "10-K"},
    ])
    d = _dilution(f, date(2026, 9, 24), price=40.0, basic_shares=400e6)
    assert abs(d["from_options"] - 10e6 * (1 - 20 / 40)) < 1 and d["from_rsus"] == 5e6 and abs(d["value"] - 410e6) < 1
    assert "treasury stock method at 40.00" in d["source"]
    assert _dilution(f, date(2026, 9, 24), price=10.0, basic_shares=400e6)["from_options"] == 0  # out of the money


def test_links_are_components_not_overrides():
    """An accepted concept link never displaces a schema total; it fills the years the total lacks and sums with other links."""
    from crucible.facts import history_from_facts
    rows = []
    def add(concept, fy, value, form="10-K", filing=None):
        end = pd.Timestamp(f"{fy}-12-31")
        rows.append({"concept": concept, "period_start": pd.Timestamp(f"{fy}-01-01"), "period_end": end, "numeric_value": value, "form_type": form,
                     "filing_date": pd.Timestamp(filing or f"{fy + 1}-02-20"), "period_type": "duration", "fiscal_period": "FY", "fiscal_year": fy,
                     "accession": f"a{fy}", "unit": "USD", "dimension": None})
    def inst(concept, fy, value):
        rows.append({"concept": concept, "period_start": pd.NaT, "period_end": pd.Timestamp(f"{fy}-12-31"), "numeric_value": value, "form_type": "10-K",
                     "filing_date": pd.Timestamp(f"{fy + 1}-02-20"), "period_type": "instant", "fiscal_period": "FY", "fiscal_year": fy,
                     "accession": f"a{fy}", "unit": "USD", "dimension": None})
    for fy in (2023, 2024, 2025):
        add("us-gaap:Revenues", fy, 1000e6)
        add("us-gaap:NetCashProvidedByUsedInOperatingActivities", fy, 100e6)
        add("us-gaap:NetCashProvidedByUsedInInvestingActivities", fy, -300e6)
        add("us-gaap:PaymentsToAcquirePropertyPlantAndEquipment", fy, 250e6)
        inst("us-gaap:Assets", fy, 5000e6)
        add("mara:PurchasedEnergyCosts", fy, 200e6)
        add("mara:HostingCosts", fy, 300e6)
        add("us-gaap:CostOfGoodsAndServicesSoldDepreciationAndAmortization", fy, 100e6)
        if fy < 2025:
            add("us-gaap:CostOfGoodsAndServicesSold", fy, 580e6)  # the total exists through 2024 only
    facts = pd.DataFrame(rows)
    links = {"mara:purchasedenergycosts": "cogs", "mara:hostingcosts": "cogs", "us-gaap:costofgoodsandservicessolddepreciationandamortization": "cogs",
             "us-gaap:paymentstoacquirepropertyplantandequipment": "cfi"}
    hist, prov = history_from_facts(facts, concept_links=links)
    assert hist.loc["cogs", 2024] == 580e6 and hist.loc["cogs", 2025] == 600e6  # total wins; components sum where it is missing
    assert hist.loc["cfi", 2025] == -300e6  # a payment linked into a cash-flow total never displaces the total
    assert any("cogs" in n and "2024" in n for n in prov["cogs"]["notes"])  # 600 vs 580 flagged
    assert "PurchasedEnergyCosts" in "".join(prov["cogs"]["concepts"])
    # rejecting the total hands the line to the components
    hist2, _ = history_from_facts(facts, concept_links=links, concept_rejects={"us-gaap:costofgoodsandservicessold"})
    assert hist2.loc["cogs", 2024] == 600e6


def test_miner_template_fit_and_build(tmp_path):
    """KPIs fitted from release text, halving-aware issuance, and a build whose identities hold."""
    from crucible.templates import annual_kpis, fit_kpis, issuance_between, miner_build, miner_defaults, subsidy_on
    rels = []
    for q, (y, m, hr, blocks, btc, cost, px, held) in enumerate([(2025, 5, 54.3, 666, 2286, 35728, 84000, 47531), (2025, 8, 57.4, 694, 2358, 33735, 98000, 49951),
                                                                  (2025, 11, 60.4, 633, 2144, 39235, 112000, 52850), (2026, 2, 66.4, 595, 2011, 48611, 96000, 53822),
                                                                  (2026, 5, 72.2, 653, 2247, 40047, 76288, 35303), (2026, 8, 70.3, 700, 2422, 38690, 71325, 35577)]):
        qn = ((m - 2) // 3) if m > 2 else 4
        yy = y if m > 2 else y - 1
        text = (f"Q{qn} {yy} results. Energized hashrate (EH/s) increased to {hr} EH/s in Q{qn} {yy}. Metric Q{qn} {yy} prior Number of Blocks Won {blocks} 600 5% "
                f"BTC Produced {btc:,} 2,000 3%. Purchased energy cost per bitcoin was ${cost:,} in Q{qn} {yy}. We produced {btc:,} BTC at an average price of ${px:,} "
                f"and sold 500 BTC at an average price of ${px + 100:,}. At quarter end, we held {held:,} bitcoin, including 9,000 pledged.")
        rels.append({"filing_date": f"{y}-{m:02d}-10", "accession": f"a{q}", "text": text})
    q = fit_kpis(rels)
    assert list(q.quarter) == ["2025Q1", "2025Q2", "2025Q3", "2025Q4", "2026Q1", "2026Q2"]
    r = q.iloc[-1]
    assert r.hashrate_eh == 70.3 and r.blocks_won == 700 and r.btc_produced == 2422 and r.energy_cost_per_btc == 38690 and r.avg_price_produced == 71325
    assert r.btc_sold == 500 and r.avg_price_sold == 71425 and r.btc_held == 35577 and r.hashrate_eh_cite.startswith("8-K EX-99.1 filed 2026-08-10")
    assert 0.04 < r.network_share < 0.07 and r.network_hashrate_eh > 1000 and abs(r.fee_share - (2422 / (700 * 3.125) - 1)) < 1e-3
    assert subsidy_on(pd.Timestamp("2024-04-19")) == 6.25 and subsidy_on(pd.Timestamp("2024-04-21")) == 3.125 and subsidy_on(pd.Timestamp("2028-06-01")) == 1.5625
    blocks, btc = issuance_between(pd.Timestamp("2027-01-01"), pd.Timestamp("2028-01-01"))
    assert abs(blocks - 52560) < 1 and abs(btc - 52560 * 3.125) < 1
    a = annual_kpis(q)
    assert a.loc[2025, "quarters"] == 4 and abs(a.loc[2025, "btc_produced"] - (2286 + 2358 + 2144 + 2011)) < 1e-6
    hist = synthetic_history()
    years = [int(hist.columns[-1]) + i for i in range(1, 6)]
    d = miner_defaults(hist, q, years)
    assert d["hashrate_last"] > 0 and d["btc_held"] == 35577 and len(d["hashrate_eh"]) == 5 and d["power_price_mwh"][0] > 0
    b = miner_build(hist, years, d)
    y0 = years[0]
    share = d["hashrate_eh"][0] / d["network_hashrate_eh"][0]
    assert abs(b.at[y0, "btc_mined"] - share * b.at[y0, "issuance_btc"] * (1 + d["fee_share"][0])) < 1e-6
    assert abs(b.at[y0, "energy_mwh"] - d["hashrate_eh"][0] * d["efficiency_j_th"][0] * 8760) < 1e-6
    assert abs(b.at[y0, "cogs"] - (b.at[y0, "energy_cost"] + b.at[y0, "other_cash_cogs"] + b.at[y0, "d_and_a"])) < 1e-6
    assert abs(b.at[y0, "capex"] - (max(d["hashrate_eh"][0] - d["hashrate_last"], 0) + d["hashrate_last"] / 3) * d["capex_per_eh"][0]) < 1e-6
    assert abs(b.at[years[-1], "treasury_value"] - 35577 * d["btc_price"][-1] / 1e6) < 1e-6  # all production sold: treasury unchanged
    # the model consumes the build through the same per-year ratios, and the workbook carries the Drivers sheet with formulas
    from crucible.excel import write_model
    from crucible.model import Drivers, WACCInputs, dcf, forecast
    from crucible.templates import apply_build_to_drivers
    drv = apply_build_to_drivers(Drivers.from_history(hist, years=5), b)
    fc = forecast(hist, drv)
    assert abs(fc.at["revenue", y0] - b.at[y0, "revenue"]) < 1e-6 and abs(fc.at["capex", y0] - b.at[y0, "capex"]) < 1e-6
    val = dcf(fc, drv, 0.10, 0.025).as_dict()
    assert abs(val["equity_value"] - (val["enterprise_value"] - val["net_debt"] + drv.non_operating_assets)) < 1e-6
    out = write_model(hist, drv, WACCInputs(0.04, 1.2, 0.05, 0.0, 0.055, 0.21, 0.2), 0.025, tmp_path / "m.xlsx", company="Synth",
                      template={"key": "bitcoin_miner", "label": "miner", "inputs": d, "build": b, "basis": d["basis"], "last_hashrate": d["hashrate_last"]})
    wb = load_workbook(out)
    assert "Drivers" in wb.sheetnames
    ws = wb["Drivers"]
    formulas = [c.value for row in ws.iter_rows(min_row=5, max_row=40) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 50 and any("8760" in f for f in formulas) and wb["DCF"]["B32"].value == "=B30-B31+B36"
    assert str(wb["Inputs"]["H9"].value).startswith("=Drivers!")  # D&A % revenue links to the driver build


def test_profile_family_quotes_and_decision_options(tmp_path):
    from crucible.profile import build_profile, family_from_sic, library_templates_for, verify_quote
    assert family_from_sic("3571")["family"] == "hardware" and family_from_sic(6199)["label"].startswith("non-bank")
    assert family_from_sic("junk", "Electronic Computers")["family"] == "generic"
    sections = {"Item 1 Business": "Item 1. Business Company Background The Company designs, manufactures and markets smartphones, personal computers and wearables. It sells services too."}
    assert verify_quote("The Company designs, manufactures and markets smartphones", sections) == "Item 1 Business"
    assert verify_quote("The Company designs teapots", sections) is None and verify_quote("too short", sections) is None
    assert library_templates_for("Electronic Computers", sections["Item 1 Business"]) == [] and library_templates_for("Finance Services", "we mine bitcoin with a fleet") == ["bitcoin_miner"]
    pk = _packet()
    prof = build_profile(pk, sections, {"ticker": "TST", "sic": "3571", "industry": "Electronic Computers"}, MockLLM())
    assert prof["family"]["family"] == "hardware" and prof["kpi_scheme_candidates"] and prof["template_candidates"][0]["value"] == "units_price" and prof["dropped"] == 0
    specs = decision_specs(sections, [], "Test Co", profile=prof)
    by = {s.key: s for s in specs}
    labels = [o.label for o in by["industry_template"].options]
    assert labels[0].endswith("(from the filing)") and any("SIC family" in l for l in labels) and "bitcoin_miner" not in labels
    assert [o.label for o in by["kpi_scheme"].options][-1] == "financial KPIs only" and "volume KPIs" in [o.label for o in by["kpi_scheme"].options]
    assert "Mock Rival Inc" in by["peer_set"].options[0].description
    # no profile: the fallbacks are industry-neutral
    plain = {s.key: s for s in decision_specs(sections, [], "Test Co")}
    assert not any("hashrate" in o.description.lower() for o in plain["kpi_scheme"].options)


def test_citation_confirmation_and_ask(tmp_path):
    from crucible.ask import ask, retrieve
    from crucible.mapsuggest import confirm_citations
    sections = {"Item 8 Notes": "Purchased energy costs consist of electricity bought for owned mining sites. Depreciation of miners is recorded in cost of revenue. Leases are immaterial.",
                "Item 7 MD&A": "Purchased energy costs rose with hashrate and power prices during the year."}
    props = [{"source": "mara:PurchasedEnergyCosts", "label": "", "statement": "IS", "target": "cogs", "relation": "energy for mining", "citations": [], "citation_auto": True},
             {"source": "Weird Line", "label": "", "statement": "IS", "target": "sga", "relation": "unclear", "citations": [], "citation_auto": True}]
    res = confirm_citations(props, sections, MockLLM())
    assert res["checked"] == 2 and res["confirmed"] == 1 and res["declined"] == 1
    assert props[0]["citation_checked"].startswith("model-confirmed") and props[0]["citations"] and not props[0]["citation_auto"]
    assert props[1]["citation_checked"].startswith("no filing sentence") and props[1]["citations"] == []
    ctx = retrieve("what drives purchased energy costs?", sections, _packet())
    assert ctx and ctx[0]["text"].startswith("Purchased energy costs") and all(c["id"].startswith("C") for c in ctx)
    a = ask("what drives purchased energy costs?", {"source": "mara:PurchasedEnergyCosts", "target": "cogs"}, _packet(), sections, MockLLM(), log_dir=tmp_path, ticker="TST")
    assert a["citations"] and a["confidence"] == 0.6 and not a["unsupported"] and (tmp_path / "qa.jsonl").exists()
    b = ask("what is the moon made of?", None, None, {"Item 8 Notes": "Nothing relevant here at all."}, MockLLM())
    assert b["unsupported"] and b["confidence"] <= 0.3


def test_review_workbook_round_trip_and_ui_service(tmp_path):
    from crucible.decisions import export_review, import_review, review_frame
    from crucible.ui import service as svc
    from openpyxl import load_workbook
    pk, spec = _packet(), _spec()
    led = DecisionLedger(tmp_path)
    led.record("TST", decide(pk, spec, MockLLM(), seed=1), pk, spec)
    led.save()
    frame = review_frame("TST", led, "decision")
    assert list(frame.columns)[-5:] == ["DECISION", "TARGET_OVERRIDE", "NOTE", "FOLLOW_UP_QUESTION", "ANSWER"] and len(frame) == 4
    path = export_review(tmp_path / "TST_review.xlsx", "TST", MappingLedger(tmp_path), led)
    wb = load_workbook(path)
    ws = wb["decisions"]
    assert ws.max_row == 5 and ws["P1"].value == "DECISION"
    ws["P2"], ws["S2"] = "accepted", "why this one?"
    ws["P3"], ws["R3"] = "unsure", "come back later"
    wb.save(path)
    res = import_review(path, "TST", MappingLedger(tmp_path), DecisionLedger(tmp_path), analyst="jc", ask_fn=lambda row, q: {"question": q, "answer": "because", "confidence": 0.7, "unsupported": False, "citations": [{"id": "C1", "source": "10-K", "text": "x"}]})
    assert res["accepted"] == 1 and res["unsure"] == 1 and res["answered"] == 1 and Path(res["answered_workbook"]).exists()
    d = DecisionLedger(tmp_path).rows("TST")
    acc = d[d.status == "accepted"].iloc[0]
    assert acc.decided_by == "jc" and "Q: why this one?" in acc.note and "conf 0.70" in acc.note
    assert svc.counts(tmp_path, "TST")["decisions"] == {"pending": 2, "accepted": 1, "unsure": 1}
    assert svc.decide(tmp_path, "TST", "decisions", "segment_scheme", d[d.status == "pending"].iloc[0].source, "rejected", note="no") == 1
    assert svc.step_args("choose", "tst", "mock")[:4] == ["choose", "TST", "--provider", "mock"]


def test_state_context_answers_model_questions(tmp_path):
    from crucible.ask import ask, retrieve
    from crucible.state import model_items, sheet_frame, sheet_names, statement_items
    hist = synthetic_history()
    from crucible.schema import MappingReport
    reports = {"IS": MappingReport(statement="IS", mapped={"revenue": "facts: Revenues"}), "BS": MappingReport(statement="BS", mapped={"cash": "facts: Cash"}), "CF": MappingReport(statement="CF")}
    items = statement_items({"hist": hist, "reports": reports, "years": list(hist.columns)}, "SYN")
    assert any(i["statement"] == "BS" and i["text"].startswith("Cash") for i in items) and all("as mapped now" in i["source"] for i in items)
    ctx = retrieve("show me what's under the balance sheet now", {"Item 8 Notes": "Cash is held in banks."}, None, state=items)
    assert ctx and ctx[0]["kind"] == "state" and all(c["kind"] == "state" for c in ctx if "balance sheet" in c["source"])
    bs_ids = [c for c in ctx if "balance sheet" in c["source"]]
    assert len(bs_ids) >= 10  # every balance-sheet line offered, not just word matches
    a = ask("show me what's under the balance sheet now", {"statement": "BS", "source": "Other Assets"}, None, {"Item 8 Notes": "Cash is held in banks."}, MockLLM(), state=items)
    assert a["citations"] and "as mapped now" in a["citations"][0]["source"] and not a["unsupported"]
    # workbook preview helpers on a freshly written model
    from crucible.excel import write_model
    from crucible.model import Drivers, WACCInputs
    out = write_model(hist, Drivers.from_history(hist, years=3), WACCInputs(0.04, 1.2, 0.05, 0.0, 0.055, 0.21, 0.2), 0.025, tmp_path / "m.xlsx", company="Synth")
    assert "Model" in sheet_names(out) and len(sheet_frame(out, "Model")) > 10
    assert model_items("NOPE", outputs=tmp_path) == []
